#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""PLC 项目配置向导 · 本地服务
静态页面 + /api/parse-table：Excel(.xlsx)/Word(.docx) IO 表解析（离线）。
仅绑定 127.0.0.1。
"""
import http.server
import io
import json
import os
import re
import socketserver

ROOT = os.path.dirname(os.path.abspath(__file__))
PORT = 8787

ADDR = re.compile(r'^%?([IQM])(\d+)\.(\d+)$', re.IGNORECASE)
SYM = re.compile(r'^[A-Za-z]{1,3}\d{1,2}[A-Za-z]?$')
CJK = re.compile(r'[一-鿿]')


def map_row(cells):
    cells = [str(c).strip() for c in cells if c is not None and str(c).strip()]
    addr, rest, sym, name = None, [], '', ''
    for c in cells:
        m = ADDR.match(c)
        if m and addr is None:
            addr = m.group(1).upper() + m.group(2) + '.' + m.group(3)
        else:
            rest.append(c)
    if addr is None:
        return None
    keep = []
    for c in rest:
        if not sym and SYM.match(c) and not CJK.search(c):
            sym = c
        else:
            keep.append(c)
    tail = []
    for c in keep:
        if not name and CJK.search(c):
            name = c.replace(' ', '')
        else:
            tail.append(c)
    if not name and tail:
        name = tail.pop(0).replace(' ', '')
    typ = 'DI' if addr[0] == 'I' else ('DO' if addr[0] == 'Q' else 'M')
    return {'address': addr, 'type': typ, 'name': name or addr,
            'symbol': sym, 'comment': ' '.join(tail)}


def parse_xlsx(data):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            r = map_row(row)
            if r:
                rows.append(r)
    return rows


def parse_docx(data):
    from docx import Document
    doc = Document(io.BytesIO(data))
    rows = []
    for table in doc.tables:
        for row in table.rows:
            r = map_row([c.text for c in row.cells])
            if r:
                rows.append(r)
    return rows


class Handler(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=ROOT, **kwargs)

    def do_POST(self):
        if self.path != '/api/parse-table':
            self.send_error(404)
            return
        n = int(self.headers.get('Content-Length', 0))
        data = self.rfile.read(n)
        fname = (self.headers.get('X-Filename') or '').lower()
        ok, rows, err = False, [], None
        try:
            if fname.endswith(('.xlsx', '.xlsm')):
                rows = parse_xlsx(data)
            elif fname.endswith('.docx'):
                rows = parse_docx(data)
            else:
                self.send_error(400, '仅支持 .xlsx / .docx')
                return
            ok = True
        except Exception as e:  # noqa: BLE001
            err = str(e)
        body = json.dumps({'ok': ok, 'rows': rows, 'error': err},
                          ensure_ascii=False).encode('utf-8')
        self.send_response(200)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, fmt, *args):
        pass


class Server(socketserver.TCPServer):
    allow_reuse_address = True


if __name__ == '__main__':
    with Server(('127.0.0.1', PORT), Handler) as httpd:
        print('PLC Wizard serving http://127.0.0.1:%d/  (xlsx/docx import)' % PORT)
        httpd.serve_forever()
